from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "/Users/akhona/ezygocouriers/deliverables/mobile-app-plan/EzyGo_Mobile_Application_Build_Plan.docx"

GREEN = "087D58"
DARK_GREEN = "153838"
AMBER = "F59E0B"
PALE_GREEN = "EEF8F2"
PALE_GRAY = "F5F7F6"
MID_GRAY = "667A74"
BORDER = "D9D9D9"


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("EzyGo Mobile Application Build Plan   |   ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="Body Text")
        indent = 0.28 if level == 0 else 0.52
        p.paragraph_format.left_indent = Inches(indent)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(4)
        p.add_run("•  ")
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph(style="Body Text")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(5)
        number = p.add_run(f"{index}.  ")
        number.bold = True
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    set_cant_split(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_fill(cell, DARK_GREEN)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)
        if widths:
            cell.width = widths[index]
    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for index, text in enumerate(row_data):
            cell = row.cells[index]
            set_cell_fill(cell, "FFFFFF" if row_index % 2 == 0 else PALE_GREEN)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text))
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string("1F2E2A")
            if widths:
                cell.width = widths[index]
    return table


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string("1F2E2A")
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12

styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(31)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].paragraph_format.space_after = Pt(14)

for level, size in ((1, 20), (2, 14), (3, 11.5)):
    style = styles[f"Heading {level}"]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = True

body = styles["Body Text"]
body.font.name = "Aptos"
body.font.size = Pt(10.5)
body.font.color.rgb = RGBColor.from_string("1F2E2A")
body.paragraph_format.space_after = Pt(8)
body.paragraph_format.line_spacing = 1.12

for style_name in ("List Bullet", "List Bullet 2", "List Number"):
    style = styles[style_name]
    style.font.name = "Aptos"
    style.font.size = Pt(10.3)
    style.font.color.rgb = RGBColor.from_string("1F2E2A")
    style.paragraph_format.space_after = Pt(4)

if "Small Text" not in styles:
    small = styles.add_style("Small Text", WD_STYLE_TYPE.PARAGRAPH)
else:
    small = styles["Small Text"]
small.font.name = "Aptos"
small.font.size = Pt(8.5)
small.font.color.rgb = RGBColor.from_string(MID_GRAY)
small.paragraph_format.space_after = Pt(5)

footer = section.footer.paragraphs[0]
add_page_number(footer)

# Cover page
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(54)
r = p.add_run("EZYGO COURIERS")
r.font.name = "Aptos"
r.font.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor.from_string(GREEN)

title = doc.add_paragraph("EzyGo Mobile Application Build Plan", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(28)
r = subtitle.add_run("Reuse assessment, target architecture and implementation roadmap")
r.font.name = "Aptos"
r.font.size = Pt(15)
r.font.color.rgb = RGBColor.from_string(MID_GRAY)

intro = doc.add_paragraph(style="Body Text")
intro.paragraph_format.space_after = Pt(16)
intro.add_run(
    "The current Next.js platform already contains the main business capabilities required by a mobile courier application. "
    "The recommended approach is to keep it as the secure backend, reuse its API and business rules, and build mobile-specific customer and driver interfaces."
)

add_heading(doc, "Recommended product direction", 2)
add_body(doc, "Start with an installable customer web app to validate demand and build a native driver app with Expo and React Native. The driver application benefits most from native background location, push notifications, navigation and offline retry. Keep the administration portal on the web.")

add_heading(doc, "What this plan covers", 2)
add_bullets(doc, [
    "The parts of the existing application that can be reused",
    "The components that must be adapted or rebuilt",
    "A recommended mobile architecture",
    "A phased implementation plan from preparation to store release",
    "Security, testing and operational readiness requirements",
])

p = doc.add_paragraph(style="Small Text")
p.paragraph_format.space_before = Pt(24)
p.add_run("Prepared from the current EzyGo Couriers repository   •   10 September 2026")

doc.add_page_break()

# Executive summary
add_heading(doc, "Executive summary", 1)
add_body(doc, "EzyGo can reuse almost all server-side business logic and data. The existing PostgreSQL database, delivery lifecycle, role checks, quote generation, Paystack integration, email notifications, driver assignment and delivery PIN logic should remain on the server. The mobile application should access them through HTTPS APIs.")
add_body(doc, "A native React Native interface cannot directly reuse the current Next.js page components because they depend on HTML, browser APIs, CSS and Next.js navigation. Shared TypeScript models, validation rules, status rules and formatting utilities can be extracted into a package used by both web and mobile.")

add_heading(doc, "Recommended delivery sequence", 2)
add_numbered(doc, [
    "Stabilise and document the existing API contracts.",
    "Implement a mobile-safe authentication and session flow.",
    "Create a shared TypeScript package for models, schemas and constants.",
    "Build the customer experience for booking, payment and tracking.",
    "Build the driver experience for assignments, status changes and background location.",
    "Add push notifications, live location controls and offline retry.",
    "Complete security testing, device testing and app-store preparation.",
])

add_heading(doc, "Recommended application split", 2)
add_table(doc,
    ["Surface", "Recommended technology", "Reason"],
    [
        ["Customer", "Responsive PWA first, then Expo if needed", "The current customer pages are already responsive and can reach users quickly."],
        ["Driver", "Expo and React Native", "Reliable background GPS, push notifications and native navigation are operational requirements."],
        ["Administrator", "Existing Next.js web portal", "Operational tables and controls are better suited to a larger web interface."],
        ["Backend", "Existing Next.js API and PostgreSQL", "The current server already owns the business rules, secrets and data."],
    ],
    [Inches(1.15), Inches(2.2), Inches(3.45)],
)

doc.add_page_break()

# Reuse assessment
add_heading(doc, "Reusable parts of the current application", 1)
add_heading(doc, "Backend services and database", 2)
add_body(doc, "The mobile application can use the current backend without duplicating core business logic. The server must remain the only component that connects to PostgreSQL or holds SMTP, OAuth and payment-provider secrets.")
add_bullets(doc, [
    "PostgreSQL schema, indexes and migrations",
    "Customer, driver and administrator roles",
    "Delivery creation, cancellation and status history",
    "Quote creation and acceptance",
    "Paystack checkout initialization, callbacks and webhooks",
    "Automatic driver assignment based on availability and recent location",
    "Delivery PIN generation and verification",
    "Authentication email and delivery-completion email",
])

add_heading(doc, "Existing customer capabilities", 2)
add_bullets(doc, [
    "Account registration, email verification and login",
    "Address selection within the Cape Town service area",
    "Pickup, recipient and parcel details",
    "Delivery creation and quote acceptance",
    "Paystack payment redirect",
    "Active delivery list and delivery history",
    "Delivery detail, status history and eligible cancellation",
])

add_heading(doc, "Existing driver capabilities", 2)
add_bullets(doc, [
    "Assigned delivery list and trip details",
    "Pickup and recipient contact information",
    "Status changes through the permitted delivery lifecycle",
    "Optional driver notes",
    "Recipient PIN confirmation for protected handovers",
    "Driver location updates and proximity-based assignment",
])

doc.add_page_break()
add_heading(doc, "Shared TypeScript candidates", 2)
add_body(doc, "Move framework-independent code into a shared package. Keep display-specific values, such as web CSS class names, outside that package.")
add_table(doc,
    ["Shared item", "Current source", "Mobile use"],
    [
        ["Delivery statuses", "lib/constants/delivery-status.ts", "Status types, labels and allowed transitions"],
        ["Delivery schema", "lib/validations/delivery.ts", "Form validation and request typing"],
        ["Service area", "lib/constants/service-area.ts", "Address guidance and local validation"],
        ["API envelope", "lib/api/response.ts", "Typed success and error handling"],
        ["Brand assets", "public directory and CSS variables", "Mobile icons, splash screens and design tokens"],
    ],
    [Inches(1.4), Inches(2.45), Inches(2.95)],
)

# Rebuild
add_heading(doc, "Parts that require mobile implementation", 1)
add_body(doc, "The existing pages provide a useful product and visual reference, but React Native uses native components rather than HTML. Recreate the interface using mobile navigation, native form controls and mobile-safe integrations.")

add_table(doc,
    ["Current web dependency", "Mobile replacement", "Reuse level"],
    [
        ["Next.js Link and routing", "Expo Router or React Navigation", "Flow and route names only"],
        ["HTML and CSS components", "React Native View, Text, Pressable and StyleSheet", "Design language and copy"],
        ["window.google Places autocomplete", "Google Places API or a maintained native component", "Cape Town bounds and address model"],
        ["window.location payment redirect", "Secure system browser with deep-link return", "Checkout endpoint and provider logic"],
        ["navigator.geolocation", "Native foreground and background location APIs", "Location endpoint and assignment service"],
        ["Browser cookie assumptions", "Secure mobile session strategy", "Server-side identity and role checks"],
    ],
    [Inches(2.15), Inches(2.85), Inches(1.8)],
)

add_heading(doc, "Capabilities to add", 2)
add_bullets(doc, [
    "Push notifications for assignments and delivery events",
    "Deep links for authentication and payment returns",
    "Offline request queue for driver status and location updates",
    "Location-permission education and foreground sharing indicator",
    "Customer-facing live driver location if this is a product requirement",
    "Crash reporting, analytics and application health monitoring",
    "Proof of delivery, such as a photograph or signature, if operations require it",
])

add_heading(doc, "Current tracking limitation", 2)
add_body(doc, "The database stores driver coordinates for assignment, but the customer API does not provide a controlled real-time driver-location stream. The existing customer tracking page mainly presents delivery status and history. Live map tracking needs a new authorized endpoint, a freshness rule and a privacy policy that limits location visibility to the correct customer and active trip.")

# Architecture
add_heading(doc, "Recommended technical architecture", 1)
add_body(doc, "Keep one business backend and treat the web and mobile products as separate clients. This prevents pricing, assignment and delivery-state rules from drifting between platforms.")

add_table(doc,
    ["Layer", "Responsibility", "Recommended implementation"],
    [
        ["Mobile interface", "Screens, navigation, device permissions and local state", "Expo and React Native with TypeScript"],
        ["Shared package", "Models, schemas, constants and API types", "Workspace package imported by web and mobile"],
        ["Mobile API client", "Base URL, authentication, retries and error mapping", "Typed fetch wrapper with request timeouts"],
        ["Backend API", "Authorization, validation and business operations", "Existing Next.js application with versioned routes"],
        ["Database", "Users, deliveries, payments, logs and driver locations", "Existing PostgreSQL database"],
        ["External services", "Payments, maps, email and push delivery", "Paystack, Google Maps, SMTP and a push provider"],
    ],
    [Inches(1.3), Inches(2.6), Inches(2.9)],
)

add_heading(doc, "Suggested repository structure", 2)
add_bullets(doc, [
    "apps/web for the existing Next.js application",
    "apps/mobile for the Expo application",
    "packages/contracts for API request and response types",
    "packages/domain for delivery statuses, validation and service-area rules",
    "packages/brand for reusable colours, spacing values and asset references",
])

add_heading(doc, "API design decisions", 2)
add_bullets(doc, [
    "Add a version prefix such as /api/v1 before releasing a mobile client that cannot be updated immediately.",
    "Return stable machine-readable error codes in addition to human-readable messages.",
    "Define request and response schemas for every mobile endpoint.",
    "Use cursor pagination for long delivery histories.",
    "Make status updates and payment confirmations idempotent so retries cannot create duplicate effects.",
    "Apply rate limits to login, OTP, location and payment endpoints.",
])

# Steps phases 1-3
add_heading(doc, "Implementation roadmap", 1)
add_heading(doc, "Phase 1 Define the mobile product", 2)
add_numbered(doc, [
    "Choose the first release audience. A sensible minimum is customers who book and track deliveries plus drivers who execute assigned trips.",
    "Write the release scope as user journeys. Include sign in, booking, payment, tracking, assignment, pickup, in transit and completion.",
    "Decide whether customer live-map tracking is required in the first release. Status tracking is already supported; live location requires additional backend and privacy work.",
    "Confirm supported platforms and operating-system versions for iOS and Android.",
    "Define operational policies for driver location, failed deliveries, cancellations and proof of delivery.",
])

add_heading(doc, "Phase 2 Prepare the backend", 2)
add_numbered(doc, [
    "Create an inventory of the existing endpoints and record each method, request body, response, authorization rule and error condition.",
    "Introduce versioned mobile API routes or guarantee compatibility for the existing routes.",
    "Extract shared delivery types, address types, API envelopes and Zod schemas into framework-independent modules.",
    "Add stable error codes, request IDs and structured server logs.",
    "Add rate limiting and abuse protection to authentication, OTP, payment and location routes.",
    "Create automated API tests covering customer ownership, driver assignment, status transitions, PIN checks and payment idempotency.",
])

add_heading(doc, "Phase 3 Implement mobile authentication", 2)
add_numbered(doc, [
    "Select a mobile session strategy supported by the authentication service. Do not place long-lived secrets in ordinary local storage.",
    "Store session material in the operating system keychain or keystore.",
    "Implement email and password sign in, registration, OTP verification, logout and session restoration.",
    "Configure application links or deep links for Google OAuth and email verification.",
    "Verify that suspended and unverified accounts cannot establish or retain sessions.",
    "Test session expiry, revoked sessions, device clock differences and network interruption.",
])

doc.add_page_break()

# phases 4-5
add_heading(doc, "Phase 4 Build the customer application", 1)
add_numbered(doc, [
    "Create the navigation shell with authentication, dashboard, new delivery, delivery detail, payment and profile routes.",
    "Build a typed API client that adds the production base URL, session credentials, timeout handling and normalized errors.",
    "Recreate the three-step booking flow with native form controls and shared validation.",
    "Integrate address autocomplete and keep the existing Cape Town service-area checks on both client and server.",
    "Submit the delivery through the existing server workflow so the quote and payment attempt are created centrally.",
    "Open Paystack in a secure browser and return through a verified app deep link.",
    "Show a payment-verification state until the backend confirms the transaction.",
    "Build active-delivery and history lists, delivery detail, status timeline and cancellation.",
    "Add pull-to-refresh, empty states, retry controls and accessible loading feedback.",
])

add_heading(doc, "Customer release screens", 2)
add_table(doc,
    ["Screen", "Required behavior", "Existing source"],
    [
        ["Authentication", "Register, OTP, sign in and session restore", "app/auth and app/api/auth"],
        ["Dashboard", "Active deliveries, history and booking entry", "app/dashboard/page.tsx"],
        ["New delivery", "Addresses, contacts, parcel options and payment", "app/dashboard/deliveries/new/page.tsx"],
        ["Delivery detail", "Status, timeline, price and cancellation", "app/dashboard/tracking/[id]/page.tsx"],
        ["Payment", "Open checkout and verify return", "app/api/payments and payment page"],
    ],
    [Inches(1.25), Inches(3.15), Inches(2.4)],
)

add_heading(doc, "Phase 5 Build the driver application", 1)
add_numbered(doc, [
    "Create driver-only navigation and enforce the driver role after session restoration.",
    "Build assignment lists and a focused current-trip screen.",
    "Show pickup, destination, parcel notes and tappable contact actions.",
    "Add navigation handoff to the device maps application.",
    "Implement status changes using the existing transition rules and endpoint.",
    "Support notes and six-digit handover PIN entry.",
    "Add foreground location first, then background tracking with explicit driver consent.",
    "Queue essential status and location requests when connectivity drops and retry them safely.",
])

# phases 6-8
add_heading(doc, "Phase 6 Add native operational features", 1)
add_heading(doc, "Push notifications", 2)
add_body(doc, "Register each device installation with the backend and associate it with the authenticated user. Send notifications from server-side delivery events rather than from another client.")
add_bullets(doc, [
    "New assignment for a driver",
    "Driver assigned for a customer",
    "Pickup confirmed",
    "Delivery completed or failed",
    "Payment confirmation or payment action required",
])

add_heading(doc, "Background location", 2)
add_body(doc, "Only track drivers when policy permits it, such as while they are active or completing a trip. Display a clear sharing state, apply battery-aware intervals and stop tracking when the driver signs out or goes off duty. The backend should reject implausible coordinates and retain only the history required for operations.")

add_heading(doc, "Live customer map", 2)
add_body(doc, "If live tracking is included, expose a reduced location payload only to the customer who owns the active delivery. Do not expose the driver’s position before assignment or after the delivery reaches a terminal state. Return a location timestamp so the app can label stale data accurately.")

add_heading(doc, "Phase 7 Test the complete system", 1)
add_table(doc,
    ["Test area", "Minimum coverage"],
    [
        ["Authentication", "Registration, OTP, login, logout, expiry, suspension and OAuth return"],
        ["Booking", "Valid and invalid addresses, phone rules, PIN email requirement and duplicate submission"],
        ["Payment", "Success, cancellation, failure, delayed webhook and repeated callback"],
        ["Driver workflow", "Assignment, pickup, transit, delivery, incorrect PIN, cancellation and failure"],
        ["Location", "Permission denied, approximate location, background mode, stale data and poor connectivity"],
        ["Security", "Role bypass, delivery ownership, sensitive-data exposure, replay and rate limits"],
        ["Devices", "Small and large phones, supported iOS and Android versions, light and dark settings"],
        ["Accessibility", "Screen reader labels, touch target size, contrast, dynamic text and keyboard behavior"],
    ],
    [Inches(1.45), Inches(5.35)],
)

add_heading(doc, "Phase 8 Prepare production release", 1)
add_numbered(doc, [
    "Create separate development, staging and production environments with distinct API URLs and provider credentials.",
    "Configure iOS and Android application identifiers, signing credentials and deep-link domains.",
    "Prepare privacy disclosures for location, personal details, notifications and analytics.",
    "Add production monitoring for API errors, payment webhooks, push delivery and background tasks.",
    "Run a controlled pilot with staff, selected drivers and selected customers.",
    "Resolve pilot issues, complete store listings and submit the production builds.",
])

# Security and decisions
add_heading(doc, "Security and privacy requirements", 1)
add_bullets(doc, [
    "Use HTTPS for every mobile request and reject production traffic over plain HTTP.",
    "Keep database, SMTP, Google OAuth and Paystack secret keys on the server.",
    "Store mobile sessions in the keychain or keystore and clear them at logout.",
    "Enforce customer ownership and driver assignment on the server for every delivery request.",
    "Keep payment webhooks authoritative and verify provider references, amounts and currency.",
    "Rate-limit authentication, OTP, payment and high-frequency location routes.",
    "Minimize driver-location retention and document when location is collected.",
    "Avoid placing recipient contact information in logs, push-notification bodies or analytics events.",
    "Provide account deletion and privacy-request procedures appropriate to the production service.",
])

add_heading(doc, "Decisions to make before development", 2)
add_table(doc,
    ["Decision", "Recommended starting position"],
    [
        ["One app or separate apps", "Use one Expo codebase with role-based navigation initially; publish separate binaries only if operations require them."],
        ["Customer PWA or native", "Launch the responsive PWA first when speed matters, while keeping the API ready for an Expo client."],
        ["Live map in first release", "Defer unless customers or operations require it; status tracking already exists."],
        ["Background driver location", "Include for the driver release, with an active-duty control and explicit permission flow."],
        ["Proof of delivery", "Confirm the operational requirement before designing uploads, storage and retention."],
        ["Offline behavior", "Queue driver status and location events; allow customer reads to fail gracefully with retry."],
    ],
    [Inches(2.0), Inches(4.8)],
)

add_heading(doc, "Suggested first release boundary", 2)
add_body(doc, "The first customer release should support account access, delivery booking, Paystack checkout, delivery lists, status detail and cancellation. The first driver release should support assignments, trip details, calling contacts, map navigation, status updates, PIN completion, push notifications and background location. Keep advanced administration, route optimization, chat, ratings and complex proof-of-delivery workflows outside the first release unless they are confirmed business requirements.")

# Checklist and endpoints
add_heading(doc, "Build readiness checklist", 1)
add_bullets(doc, [
    "The production backend has a stable HTTPS domain.",
    "Every mobile endpoint has a documented request and response contract.",
    "Mobile authentication has been selected and tested on physical devices.",
    "Deep links work from email, OAuth and Paystack on iOS and Android.",
    "Shared schemas and types compile in both the web and mobile projects.",
    "Driver location policies and permission copy have been approved.",
    "Push-notification tokens can be registered, refreshed and revoked.",
    "Payment completion remains correct when callbacks arrive more than once.",
    "The application has staging data and test users for each role.",
    "Monitoring covers API failures, payment webhooks and background location.",
    "Privacy policy and store disclosures match actual data collection.",
    "A pilot group and support process are ready before public release.",
])

add_heading(doc, "Current API foundation", 2)
add_table(doc,
    ["Area", "Current route", "Mobile purpose"],
    [
        ["Current user", "GET /api/auth/me", "Restore identity and role"],
        ["Login", "POST /api/auth/login", "Establish an authenticated session"],
        ["Customer deliveries", "GET and POST /api/deliveries", "List and create deliveries"],
        ["Delivery detail", "GET /api/deliveries/[id]", "Show delivery and status logs"],
        ["Customer action", "POST /api/deliveries/[id]", "Confirm or cancel a delivery"],
        ["Create payment", "POST /api/payments/create", "Start Paystack checkout"],
        ["Driver deliveries", "GET /api/driver/deliveries", "List assignments"],
        ["Driver detail", "GET /api/driver/deliveries/[id]", "Show trip details and logs"],
        ["Driver status", "PATCH /api/driver/status", "Update the delivery lifecycle"],
        ["Driver location", "PATCH /api/driver/location", "Send coordinates and check the queue"],
    ],
    [Inches(1.45), Inches(2.6), Inches(2.75)],
)

add_heading(doc, "Immediate next actions", 2)
add_numbered(doc, [
    "Approve the first release boundary and decide whether the customer product starts as a PWA or Expo app.",
    "Create an API contract document and mobile-safe authentication proof of concept.",
    "Restructure the repository into web, mobile and shared packages without changing production behavior.",
    "Build a thin vertical slice: sign in, load deliveries, open one delivery and sign out.",
    "After the slice works on physical iOS and Android devices, implement booking and the complete driver trip workflow.",
])

add_heading(doc, "Key repository references", 2)
add_table(doc,
    ["Current source", "Why it matters for mobile"],
    [
        ["app/api/deliveries/route.ts", "Customer delivery list and booking entry point"],
        ["app/api/driver/status/route.ts", "Driver status update contract"],
        ["app/api/driver/location/route.ts", "Driver coordinate upload contract"],
        ["app/api/payments/create/route.ts", "Paystack checkout initialization"],
        ["lib/auth/session.ts", "Current cookie-based session boundary"],
        ["lib/constants/delivery-status.ts", "Reusable delivery lifecycle rules"],
        ["lib/validations/delivery.ts", "Reusable booking validation schema"],
        ["components/driver/DriverLocationTracker.tsx", "Browser tracker to replace with native background location"],
    ],
    [Inches(3.1), Inches(3.7)],
)

doc.core_properties.title = "EzyGo Mobile Application Build Plan"
doc.core_properties.subject = "Reuse assessment and phased implementation roadmap"
doc.core_properties.author = "EzyGo Couriers"
doc.core_properties.keywords = "EzyGo, mobile, Expo, React Native, courier, implementation plan"
doc.save(OUTPUT)
print(OUTPUT)
